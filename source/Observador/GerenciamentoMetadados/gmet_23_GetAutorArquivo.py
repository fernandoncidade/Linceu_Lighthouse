import os
from utils.LogManager import LogManager
logger = LogManager.get_logger()

def get_autor_arquivo(item, loc):
    caminho = item.get("dir_atual") or item.get("dir_anterior")
    if not caminho or not os.path.exists(caminho):
        if "autor" in item and item["autor"]:
            return loc.traduzir_metadados(item["autor"], "autor")

        return ""

    autor = ""
    base, ext = os.path.splitext(caminho)
    ext = ext.lower()

    try:
        if os.path.isdir(caminho):
            try:
                import win32security
                sd = win32security.GetFileSecurity(caminho, win32security.OWNER_SECURITY_INFORMATION)
                owner_sid = sd.GetSecurityDescriptorOwner()

                try:
                    nome, dominio, tipo = win32security.LookupAccountSid(None, owner_sid)
                    autor = f"{dominio}\\{nome}"

                except Exception as lookup_err:
                    logger.warning(f"Não foi possível resolver SID para nome: {lookup_err}")
                    try:
                        sid_str = win32security.ConvertSidToStringSid(owner_sid)
                        autor = sid_str

                    except Exception as conv_err:
                        logger.error(f"Erro ao converter SID para string: {conv_err}", exc_info=True)
                        autor = ""

            except Exception as e:
                logger.error(f"Erro ao obter proprietário da pasta: {e}", exc_info=True)
                autor = ""

        else:
            if ext in [".docx", ".dotx", ".docm", ".dotm"]:
                from docx import Document
                doc = Document(caminho)
                props = doc.core_properties
                autor = (props.author or "").strip()

            elif ext == ".doc":
                import olefile
                try:
                    if olefile.isOleFile(caminho):
                        with olefile.OleFileIO(caminho) as ole:
                            if ole.exists('\x05SummaryInformation'):
                                props = ole.getproperties('\x05SummaryInformation')
                                autor = props.get(4, "") or ""
                                if isinstance(autor, bytes):
                                    autor = autor.decode("latin-1", errors="ignore")

                                autor = autor.strip()

                    else:
                        autor = ""

                except Exception as e:
                    logger.error(f"Erro ao obter autor de {ext} usando olefile: {e}", exc_info=True)
                    autor = ""

            elif ext in [".xlsx", ".xlsm", ".xltx", ".xltm"]:
                import zipfile
                try:
                    if zipfile.is_zipfile(caminho):
                        try:
                            from openpyxl import load_workbook
                            wb = load_workbook(caminho, read_only=True, data_only=True)
                            if wb.properties.creator:
                                autor = str(wb.properties.creator).strip()

                            wb.close()

                        except Exception as xlsx_err:
                            try:
                                with zipfile.ZipFile(caminho, 'r') as z:
                                    if 'docProps/core.xml' in z.namelist():
                                        import xml.etree.ElementTree as ET
                                        core_xml = z.read('docProps/core.xml')
                                        root = ET.fromstring(core_xml)
                                        ns = {
                                            'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
                                            'dc': 'http://purl.org/dc/elements/1.1/',
                                            'dcterms': 'http://purl.org/dc/terms/'
                                        }
                                        node = root.find('.//dc:creator', ns)
                                        autor = (node.text or "").strip() if node is not None else ""

                            except Exception as zerr:
                                logger.error(f"Fallback core.xml falhou: {zerr}", exc_info=True)
                                autor = ""

                    else:
                        autor = ""

                except Exception as xlsx_err:
                    logger.error(f"Erro ao ler XLSX com openpyxl: {xlsx_err}", exc_info=True)
                    autor = ""

            elif ext == ".xls":
                import olefile
                try:
                    if olefile.isOleFile(caminho):
                        with olefile.OleFileIO(caminho) as ole:
                            if ole.exists('\x05SummaryInformation'):
                                props = ole.getproperties('\x05SummaryInformation')
                                autor = props.get(4, "") or ""
                                if isinstance(autor, bytes):
                                    autor = autor.decode("latin-1", errors="ignore")

                                autor = autor.strip()

                    else:
                        import zipfile, xml.etree.ElementTree as ET
                        if zipfile.is_zipfile(caminho):
                            try:
                                with zipfile.ZipFile(caminho, 'r') as z:
                                    if 'docProps/core.xml' in z.namelist():
                                        core_xml = z.read('docProps/core.xml')
                                        root = ET.fromstring(core_xml)
                                        ns = {'dc': 'http://purl.org/dc/elements/1.1/'}
                                        node = root.find('.//dc:creator', ns)
                                        autor = (node.text or "").strip() if node is not None else ""

                            except Exception as e2:
                                logger.error(f"Erro ao extrair autor via core.xml: {e2}", exc_info=True)
                                autor = ""

                        else:
                            autor = ""

                except Exception as e:
                    logger.error(f"Erro ao obter autor de {ext} usando olefile: {e}", exc_info=True)
                    autor = ""

            elif ext in [".pptx", ".potx", ".ppsx"]:
                from pptx import Presentation
                pres = Presentation(caminho)
                autor = (pres.core_properties.author or "").strip()

            elif ext == ".ppt":
                import olefile
                try:
                    if olefile.isOleFile(caminho):
                        with olefile.OleFileIO(caminho) as ole:
                            if ole.exists('\x05SummaryInformation'):
                                props = ole.getproperties('\x05SummaryInformation')
                                autor = props.get(4, "") or ""
                                if isinstance(autor, bytes):
                                    autor = autor.decode("latin-1", errors="ignore")

                                autor = autor.strip()

                    else:
                        autor = ""

                except Exception as e:
                    logger.error(f"Erro ao obter autor de {ext} usando olefile: {e}", exc_info=True)
                    autor = ""

            elif ext in [".mdb", ".accdb"]:
                autor = ""

            elif ext == ".msg":
                import olefile
                try:
                    if olefile.isOleFile(caminho):
                        with olefile.OleFileIO(caminho) as ole:
                            if ole.exists('\x05SummaryInformation'):
                                props = ole.getproperties('\x05SummaryInformation')
                                autor = props.get(4, "") or ""
                                if isinstance(autor, bytes):
                                    autor = autor.decode("latin-1", errors="ignore")

                                autor = autor.strip()

                            elif ole.exists('__properties_version1.0'):
                                props = ole.getproperties('__properties_version1.0')
                                autor_val = None
                                for prop_id in [0x0C1A, 0x0E04, 0x0042]:
                                    if prop_id in props:
                                        autor_val = props[prop_id]
                                        break

                                if isinstance(autor_val, bytes):
                                    autor_val = autor_val.decode("latin-1", errors="ignore")

                                autor = (autor_val or "").strip()

                            else:
                                autor = ""

                    else:
                        autor = ""

                except Exception as e:
                    logger.error(f"Erro ao extrair informações do MSG: {e}", exc_info=True)
                    autor = ""

            elif ext in [".pst", ".ost"]:
                autor = ""

            elif ext == ".pub":
                autor = ""

            elif ext in [".vsd", ".vsdx"]:
                autor = ""

            elif ext in [".mpp", ".mpt"]:
                autor = ""

            elif ext == ".pdf":
                from PyPDF2 import PdfReader
                reader = PdfReader(caminho)
                info = reader.metadata
                autor = ((info.author or "") if info else "").strip()

            elif ext in [".txt", ".htm", ".html", ".mht", ".mhtml"]:
                autor = ""

            elif ext in [".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz", ".cab"]:
                autor = ""

            else:
                autor = ""

    except Exception as e:
        logger.error(f"Erro ao obter autor do arquivo {caminho}: {e}", exc_info=True)
        pass

    return loc.traduzir_metadados(autor, "autor")
